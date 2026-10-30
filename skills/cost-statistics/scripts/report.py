"""Create a compact editable Word report from aggregate.py output."""
import argparse
import datetime as dt
import json
import math
import os
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from aggregate import LABELS,InputError,finite


def duration(value):
    if value is None:return '未取得'
    n=math.floor(value+.5);h,r=divmod(n,3600);m,s=divmod(r,60)
    return f'{h}时{m:02d}分{s:02d}秒'


def shown(value,kind):
    if value is None:return '未取得'
    if kind=='time':return duration(value)
    if kind=='token':return f'{value/10000:.4f}'
    if kind=='rate':return f'{value:.2%}'
    return f'{value:.2f}'


def font_path(explicit=None):
    choices=[explicit,os.environ.get('COST_STATISTICS_FONT'),
        'C:/Windows/Fonts/msyh.ttc','C:/Windows/Fonts/simhei.ttf',
        '/System/Library/Fonts/PingFang.ttc','/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc']
    for item in choices:
        if item and Path(item).is_file():return str(item)
    raise InputError('未找到中文字体，请用--font指定本机中文字体文件')


def relative(a,b):
    return (a/b-1) if a is not None and b and finite(a) and finite(b) else None


def suggested_analysis(data):
    """Deterministic numerical observations; the executing agent may refine them."""
    objects=data['objects'];focus=next(o for o in objects if o['name']==data['focus'])
    others=[o for o in objects if o is not focus]
    other=next((o for o in others if o['name']=='WorkBuddy'),others[0])
    a,b=focus['means'],other['means'];notes=[];name=focus['name'];foil=other['name']
    if all(finite(v) and v>0 for v in (focus['input_per_call'],other['input_per_call'],a['llm_calls'],b['llm_calls'],a['total_tokens'],b['total_tokens'])):
        change=relative(a['total_tokens'],b['total_tokens'])
        direction='高' if change>0 else '低' if change<0 else '相同'
        comparison=f'{direction}{abs(change):.1%}' if change else '相同'
        text=f'{name}的平均总Token与{foil}相比{comparison}。按累计输入量除以调用数，单次LLM调用平均输入分别为{focus["input_per_call"]/10000:.2f}万Token和{other["input_per_call"]/10000:.2f}万Token。'
        calls=relative(a['llm_calls'],b['llm_calls']);inputs=relative(a['input_tokens'],b['input_tokens'])
        if calls is not None and inputs is not None and calls*inputs<0:
            text+=f'{name}的调用次数'+('更多，但累计输入量较少，差异主要体现在较小的单次输入规模。' if calls>0 else '更少，但累计输入量较多，差异主要体现在较大的单次输入规模。')
        elif calls is not None and inputs is not None:
            text+='累计输入量同时受到调用次数和单次输入规模的影响。'
        notes.append(dict(heading='调用规模与Token消耗',text=text))
    if all(finite(a[k]) for k in ('p50','p95','seconds')):
        text=f'{name}的端到端耗时P50为{duration(a["p50"])}，P95为{duration(a["p95"])}。'
        if all(finite(b[k]) for k in ('p50','p95')) and (a['p50']-b['p50'])*(a['p95']-b['p95'])<0:
            text+=f'其与{foil}在中位数和高分位处排序不同，耗时差异随运行任务而变化。'
        if a['seconds']>a['p50']:
            text+='平均耗时高于中位数，较长运行抬高了整体均值。'
        elif a['seconds']<a['p50']:
            text+='平均耗时低于中位数，较短运行拉低了整体均值。'
        notes.append(dict(heading='运行耗时分布',text=text))
    rates=[o for o in objects if finite(o['tool_error_rate'])]
    if rates:
        text='按累计工具错误次数除以累计工具调用数，错误占比为'+'、'.join(o['name']+f'{o["tool_error_rate"]:.2%}' for o in rates)+'。'
        if len(rates)>1:
            lo=min(rates,key=lambda o:o['tool_error_rate']);hi=max(rates,key=lambda o:o['tool_error_rate'])
            if lo['tool_error_rate']<hi['tool_error_rate']:text+=lo['name']+'的错误占比较低，'+hi['name']+'的错误占比较高。'
        notes.append(dict(heading='工具调用与错误频度',text=text))
    return notes[:3]


def make_chart(data,target,fontfile):
    objects=data['objects'];n=len(objects);height=210+n*94
    image=Image.new('RGB',(1700,height),'white');draw=ImageDraw.Draw(image)
    font=lambda size:ImageFont.truetype(fontfile,size)
    palette=['#173F65','#168F83','#7D91A5','#A36E32','#746490','#587F64']
    for left,title,key,factor in [(30,'平均总Token（万）','total_tokens',10000),(875,'平均端到端耗时','seconds',60)]:
        seconds_axis=key=='seconds' and max((o['means'][key] for o in objects if o['means'][key] is not None),default=0)<120
        if seconds_axis:factor=1
        draw.text((left,18),title,font=font(34),fill='#163F64')
        values=[o['means'][key]/factor if o['means'][key] is not None else None for o in objects]
        largest=max((v for v in values if v is not None),default=0)
        if largest:
            raw_step=largest*1.1/3;power=10**math.floor(math.log10(raw_step))
            step=next(v*power for v in (1,1.5,2,2.5,3,5,10) if v*power>=raw_step)
            maximum=step*3
        else:maximum=1;step=1/3
        precision=max(0,-math.floor(math.log10(step)))
        if not math.isclose(step,round(step,precision)):precision+=1
        x0=left+190;width=425;top=115;bottom=top+(n-1)*94+70
        for tick in range(4):
            x=x0+width*tick/3;draw.line((x,top-15,x,bottom),fill='#E7EDF2',width=2)
            v=maximum*tick/3;label=f'{v:.{precision}f}'
            if key=='seconds':label+='秒' if seconds_axis else '分'
            draw.text((x-12,bottom+14),label,font=font(25),fill='#62727E')
        for i,(obj,v) in enumerate(zip(objects,values)):
            y=top+i*94;name=obj['name'];namefont=font(28)
            while draw.textlength(name,font=namefont)>180 and namefont.size>17:namefont=font(namefont.size-1)
            draw.text((left,y+9),name,font=namefont,fill='#263645')
            if v is not None:
                end=x0+width*v/maximum;draw.rectangle((x0,y,end,y+45),fill=palette[i%len(palette)])
                label=(f'{v:.4f}' if abs(v)<1 else f'{v:.2f}') if key=='total_tokens' else duration(obj['means'][key])
                draw.text((end+10,y+10),label,font=font(26),fill='#263645')
            else:draw.text((x0+10,y+10),'未取得',font=font(26),fill='#62727E')
    draw.text((30,height-40),f'共同题目{len(data["common_tasks"])}道  各题等权',font=font(23),fill='#647481')
    image.save(target)


def add_table(doc,rows,widths,header=True):
    table=doc.add_table(rows=0,cols=len(widths));table.autofit=False
    for column,w in zip(table.columns,widths):column.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        element=OxmlElement('w:'+edge)
        for attr,v in [('val','single'),('sz','4'),('color','D9D9D9')]:element.set(qn('w:'+attr),v)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        row=table.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if not i and header:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,(cell,text,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin=OxmlElement('w:tcMar')
            for edge,amount in [('top',100),('bottom',100),('left',100),('right',100)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(amount));e.set(qn('w:type'),'dxa');margin.append(e)
            cell._tc.get_or_add_tcPr().append(margin)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(str(text));run.font.size=Pt(10)
            fill='173F65' if i==0 and header else 'F0F4F7' if header and i%2==0 else None
            if fill:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(e)
            if i==0 and header:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    return table


def build(data,target,notes=None,date=None,font=None):
    target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
    work=target.parent/'.cost-statistics';work.mkdir(exist_ok=True)
    fontfile=font_path(font);family=ImageFont.truetype(fontfile,12).getname()[0]
    chart=work/'comparison.png';make_chart(data,chart,fontfile)
    objects=data['objects'];notes=suggested_analysis(data) if notes is None else notes
    if len(notes)>3 or any(not n.get('heading') or not n.get('text') or len(n['text'])>260 for n in notes):
        raise InputError('分析应为最多3段，每段含heading、text，正文每段不超过260字')
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.9);sec.left_margin=sec.right_margin=Cm(2)
    sec.footer_distance=Cm(.8);sec.different_first_page_header_footer=True
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        style=doc.styles[name];style.font.name=family;style.font.color.rgb=RGBColor(0,0,0);style.font.italic=False
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family);style.paragraph_format.space_after=Pt(7)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.3
    doc.styles['Title'].font.size=Pt(36);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.color.rgb=RGBColor.from_string('5A6B78')
    def paragraph(text='',style=None):
        p=doc.add_paragraph(text,style);p.paragraph_format.keep_together=True;return p
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('成本统计报告  |  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('647481')
    paragraph().paragraph_format.space_after=Pt(70)
    paragraph('成本统计报告','Title');paragraph('执行过程消耗对比','Subtitle')
    p=paragraph(' / '.join(o['name'] for o in objects));p.paragraph_format.space_before=Pt(22)
    for run in p.runs:run.font.size=Pt(14);run.font.color.rgb=RGBColor.from_string('173F65')
    paragraph().paragraph_format.space_after=Pt(28)
    sources=sorted({t['source'] for t in data['common_tasks']})
    rows=[['受测对象','、'.join(o['name'] for o in objects)],['任务范围','、'.join(sources)],
        ['样本规模',f'共同题目{len(data["common_tasks"])}道，共{sum(o["run_count"] for o in objects)}次运行'],
        ['运行次数','；'.join(o['name']+str(o['run_count'])+'次' for o in objects)]]
    dates=sorted({v for o in objects for v in o['dates']})
    if dates:rows.append(['测试日期',dates[0] if len(dates)==1 else dates[0]+' 至 '+dates[-1]])
    rows.append(['报告日期',date or dt.date.today().isoformat()]);add_table(doc,rows,[3,14],False)
    paragraph('比较任务执行的Token消耗、耗时及调用情况。').paragraph_format.space_before=Pt(22)
    if len(objects)>3:
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
        sec.different_first_page_header_footer=False;width=25.7
    else:doc.add_page_break();width=17
    paragraph('执行消耗统计','Heading 1')
    paragraph(f'比较{len(objects)}个对象在{len(data["common_tasks"])}道共同题目上的运行消耗，各项平均指标按题目等权汇总。')
    paragraph('表1  共同任务执行指标对比','Caption').paragraph_format.keep_with_next=True
    rows=[['指标',*[o['name'] for o in objects]]]
    for label,key,kind in LABELS:rows.append([label,*[shown(o['means'][key],kind) for o in objects]])
    add_table(doc,rows,[6.2]+[(width-6.2)/len(objects)]*len(objects))
    paragraph('指标说明','Heading 2').paragraph_format.space_before=Pt(14)
    paragraph('平均指标先求同题各次运行均值，再按题目等权汇总。总Token为输入与输出之和；缓存命中率采用同题均值的算术平均。LLM完成调用数统计已完成的模型调用；工具错误统计明确错误状态或非零退出码的调用。')
    paragraph('端到端耗时为用户提交至关联执行链最后完成回复的时长。P50、P95为逐次运行耗时的50%、95%分位数，按Excel PERCENTILE.INC线性插值计算。')
    incomplete={k:v for k,v in data['metric_task_counts'].items() if v!=len(data['common_tasks'])}
    if incomplete:
        labels={k:label for label,k,_ in LABELS}
        paragraph('共同有效题数：'+'；'.join(labels[k]+' '+str(v)+'题' for k,v in incomplete.items())+'。')
    doc.add_page_break();paragraph('消耗差异分析','Heading 1')
    paragraph('图1  平均总Token与平均端到端耗时对比','Caption').paragraph_format.keep_with_next=True
    p=paragraph();p.add_run().add_picture(str(chart),width=Cm(min(width,19)))
    p.paragraph_format.space_after=Pt(16)
    for item in notes:
        paragraph(item['heading'],'Heading 2');paragraph(item['text'])
    doc.core_properties.title='成本统计报告';doc.core_properties.author='';doc.core_properties.subject='受测对象执行过程消耗对比'
    doc.save(target);return target


def main():
    p=argparse.ArgumentParser(description='从已汇总的成本数据生成Word报告')
    p.add_argument('data');p.add_argument('--output',required=True);p.add_argument('--analysis',help='[{heading,text}]，由执行Agent依据汇总数据撰写')
    p.add_argument('--date');p.add_argument('--font');args=p.parse_args()
    data=json.loads(Path(args.data).read_text(encoding='utf-8-sig'))
    notes=json.loads(Path(args.analysis).read_text(encoding='utf-8-sig')) if args.analysis else None
    print(build(data,args.output,notes,args.date,args.font).resolve())


if __name__=='__main__':
    try:main()
    except (InputError,OSError,KeyError) as exc:raise SystemExit('报告生成停止：'+str(exc))
